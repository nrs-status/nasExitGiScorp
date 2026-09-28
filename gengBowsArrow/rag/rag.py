"""rag: retrieval-augmented question answering over local documents.

Pipeline (each stage's model is set in the TOML config):
  vision model    -> describes/transcribes images (and scanned pages)
  embedding model -> embeds document chunks and the question
  chat model      -> answers the question from the retrieved chunks

usage: rag [--config FILE.toml] [--api-key-file P] [--chat-model M] [question ...]

The config file path comes from --config, else $RAG_CONFIG (the command
line option has the higher precedence); one of them is required.

Config example:
  api_key_file = "/run/secrets/keys/openrouter"
  [models]
  chat = "openai/gpt-4o-mini"
  embedding = "openai/text-embedding-3-small"
  vision = "openai/gpt-4o-mini"
  [documents]
  paths = ["docs", "report.pdf"]     # files or directories (recursive)
  [agent]                            # all optional
  top_k = 6
  chunk_size = 1200
  chunk_overlap = 200
  system_prompt = "..."
  api_url = "https://openrouter.ai/api/v1"
  ipv4_only = true                   # default true (ISP blackholes IPv6); false allows IPv6
Relative paths are resolved against the config file's directory.
"""
import argparse
import base64
import io
import math
import os
import subprocess
import sys
import tempfile
import tomllib

import requests

DJVUTXT = "@DJVUTXT@"
DDJVU = "@DDJVU@"

TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json",
            ".yaml", ".yml", ".toml", ".xml", ".html", ".htm", ".log",
            ".ini", ".cfg", ".tex", ".org"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif",
             ".tiff"}
SHEET_EXT = {".xlsx", ".xlsm", ".xls"}
DEFAULT_SYSTEM = ("You answer questions using only the provided document "
                  "excerpts. Cite sources as [file]. If the excerpts do not "
                  "contain the answer, say so.")


def log(*a):
    print(*a, file=sys.stderr)


class Api:
    def __init__(self, url, key):
        self.url = url.rstrip("/")
        self.h = {"Authorization": "Bearer " + key}

    def post(self, path, body):
        r = requests.post(self.url + path, json=body, headers=self.h,
                          timeout=(15, 180))
        if r.status_code != 200:
            raise RuntimeError(f"{path}: HTTP {r.status_code}: {r.text[:300]}")
        return r.json()

    def chat(self, model, messages):
        j = self.post("/chat/completions",
                      {"model": model, "messages": messages})
        return j["choices"][0]["message"]["content"] or ""

    def embed(self, model, texts):
        out = []
        for i in range(0, len(texts), 64):
            j = self.post("/embeddings",
                          {"model": model, "input": texts[i:i + 64]})
            data = sorted(j["data"], key=lambda d: d["index"])
            out += [d["embedding"] for d in data]
        return out


def describe_image(api, model, img_bytes, name):
    """Return text for an image (any PIL-readable format; first frame)."""
    from PIL import Image
    im = Image.open(io.BytesIO(img_bytes))
    im.seek(0)
    im = im.convert("RGB")
    im.thumbnail((2000, 2000))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    url = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    return api.chat(model, [{"role": "user", "content": [
        {"type": "text", "text": "Transcribe all text in this image and "
         "describe its content in detail so it can be searched later."},
        {"type": "image_url", "image_url": {"url": url}}]}])


def read_pdf(path, api, vision):
    import logging
    from PIL import Image
    from pypdf import PdfReader
    logging.getLogger("pypdf").setLevel(logging.ERROR)  # quiet parser noise
    rd = PdfReader(path)
    if rd.is_encrypted and not rd.decrypt(""):
        raise ValueError("PDF is password protected")
    parts = []
    for n, page in enumerate(rd.pages, 1):
        t = (page.extract_text() or "").strip()
        if t:
            parts.append(f"[page {n}]\n{t}")
        # describe images on scanned pages, and large images on text pages
        try:
            imgs = list(page.images)[:3]
        except Exception as e:
            log(f"warn: {path} p{n}: {e}")
            continue
        for im in imgs:
            try:
                if len(t) >= 50:
                    w, h = Image.open(io.BytesIO(im.data)).size
                    if w < 300 or h < 300:
                        continue
                parts.append(f"[page {n} image]\n" + describe_image(
                    api, vision, im.data, path))
            except Exception as e:
                log(f"warn: {path} p{n}: {e}")
    return "\n\n".join(parts)


def read_djvu(path, api, vision):
    r = subprocess.run([DJVUTXT, path], capture_output=True, text=True,
                       timeout=120)
    if r.stdout.strip():
        return r.stdout
    # no text layer: render first pages and send to the vision model
    parts = []
    with tempfile.TemporaryDirectory() as d:
        for n in range(1, 11):
            out = os.path.join(d, f"{n}.tif")
            p = subprocess.run([DDJVU, "-format=tiff", f"-page={n}", path,
                                out], capture_output=True, timeout=120)
            if p.returncode != 0 or not os.path.exists(out):
                break
            with open(out, "rb") as f:
                parts.append(f"[page {n}]\n" + describe_image(
                    api, vision, f.read(), path))
    return "\n\n".join(parts)


def read_docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def read_sheet(path):
    rows_out = []
    if path.lower().endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(path)
        for s in wb.sheets():
            rows_out.append(f"[sheet {s.name}]")
            for r in range(s.nrows):
                rows_out.append(" | ".join(str(c.value) for c in s.row(r)))
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for s in wb.worksheets:
            rows_out.append(f"[sheet {s.title}]")
            for row in s.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    rows_out.append(" | ".join(
                        "" if c is None else str(c) for c in row))
    return "\n".join(rows_out)


def load(path, api, models):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return read_pdf(path, api, models["vision"])
    if ext in (".djvu", ".djv"):
        return read_djvu(path, api, models["vision"])
    if ext == ".docx":
        return read_docx(path)
    if ext in SHEET_EXT:
        return read_sheet(path)
    if ext in IMAGE_EXT:
        with open(path, "rb") as f:
            return describe_image(api, models["vision"], f.read(), path)
    if ext in TEXT_EXT or ext == "":
        with open(path, encoding="utf-8", errors="replace") as f:
            return f.read()
    raise ValueError("unsupported file type " + ext)


def expand(paths, base):
    for p in paths:
        p = os.path.join(base, os.path.expanduser(p))
        if os.path.isdir(p):
            for root, _, files in sorted(os.walk(p)):
                for f in sorted(files):
                    yield os.path.join(root, f)
        else:
            yield p


def chunk(text, size, overlap):
    text = text.strip()
    step = max(1, size - overlap)
    return [text[i:i + size] for i in range(0, len(text), step)
            if text[i:i + size].strip()]


def cosine(a, b):
    d = sum(x * y for x, y in zip(a, b))
    n = math.sqrt(sum(x * x for x in a)) * math.sqrt(sum(y * y for y in b))
    return d / n if n else 0.0


def answer(api, cfg, index, question):
    qv = api.embed(cfg["models"]["embedding"], [question])[0]
    top = sorted(index, key=lambda c: -cosine(qv, c[2]))[:cfg["top_k"]]
    ctx = "\n\n".join(f"[{os.path.basename(f)}]\n{t}" for f, t, _ in top)
    return api.chat(cfg["models"]["chat"], [
        {"role": "system", "content": cfg["system_prompt"]},
        {"role": "user", "content": f"Excerpts:\n{ctx}\n\nQuestion: {question}"}])


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--config", default=os.environ.get("RAG_CONFIG"),
                    help="TOML config file (default: $RAG_CONFIG)")
    ap.add_argument("--api-key-file")
    ap.add_argument("--chat-model")
    ap.add_argument("question", nargs="*")
    a = ap.parse_args()
    if not a.config:
        ap.error("no config file: pass --config or set $RAG_CONFIG")

    base = os.path.dirname(os.path.abspath(a.config))
    try:
        with open(a.config, "rb") as f:
            c = tomllib.load(f)
    except (OSError, tomllib.TOMLDecodeError) as e:
        sys.exit(f"rag: cannot read config {a.config}: {e}")
    ag = c.get("agent", {})
    cfg = {
        "models": {"chat": "openai/gpt-4o-mini",
                   "embedding": "openai/text-embedding-3-small",
                   "vision": "openai/gpt-4o-mini", **c.get("models", {})},
        "top_k": ag.get("top_k", 6),
        "system_prompt": ag.get("system_prompt", DEFAULT_SYSTEM),
    }
    if a.chat_model:
        cfg["models"]["chat"] = a.chat_model
    keyfile = a.api_key_file or c.get("api_key_file") or \
        "/run/secrets/keys/openrouter"
    keyfile = os.path.join(base, os.path.expanduser(keyfile))
    with open(keyfile) as f:
        key = f.read().strip()
    if ag.get("ipv4_only", True):
        import socket
        import urllib3.util.connection as uc
        uc.allowed_gai_family = lambda: socket.AF_INET
    api = Api(ag.get("api_url", "https://openrouter.ai/api/v1"), key)

    size = ag.get("chunk_size", 1200)
    ov = ag.get("chunk_overlap", 200)
    chunks = []
    for p in expand(c.get("documents", {}).get("paths", []), base):
        try:
            text = load(p, api, cfg["models"])
        except Exception as e:
            log(f"skip {p}: {e}")
            continue
        cs = chunk(text, size, ov)
        log(f"loaded {p}: {len(cs)} chunk(s)")
        chunks += [(p, t) for t in cs]
    if not chunks:
        sys.exit("no document content could be loaded")
    vecs = api.embed(cfg["models"]["embedding"], [t for _, t in chunks])
    index = [(f, t, v) for (f, t), v in zip(chunks, vecs)]

    if a.question:
        print(answer(api, cfg, index, " ".join(a.question)))
        return
    try:
        while True:
            q = input("> ").strip()
            if q:
                print(answer(api, cfg, index, q))
    except EOFError:
        pass


if __name__ == "__main__":
    main()
